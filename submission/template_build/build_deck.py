"""Fill the official AI Builder Cup template with Rakshak content (team Prod Pilots).

Usage: python build_deck.py [--firebase-hosting URL]
Writes ../Rakshak_Pitch_Deck_Prod_Pilots.pptx
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

import pypdfium2 as pdfium
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "shots"
OUT = HERE.parent / "Rakshak_Pitch_Deck_Prod_Pilots.pptx"
HOSTING = None
if "--firebase-hosting" in sys.argv:
    HOSTING = sys.argv[sys.argv.index("--firebase-hosting") + 1]

URL = "https://rakshak-223197895518.asia-southeast1.run.app"
REPO = "https://github.com/devaramanelohith/rakshak-agent"

INK = RGBColor(0x20, 0x21, 0x24)
MUTED = RGBColor(0x5F, 0x63, 0x68)
BLUE = RGBColor(0x1A, 0x73, 0xE8)
RED = RGBColor(0xC5, 0x22, 0x1F)
GREEN = RGBColor(0x18, 0x80, 0x38)
AMBER = RGBColor(0xB0, 0x60, 0x00)
CARD = RGBColor(0xF1, 0xF3, 0xF4)
BLUE_T = RGBColor(0xE8, 0xF0, 0xFE)
RED_T = RGBColor(0xFC, 0xE8, 0xE6)
GREEN_T = RGBColor(0xE6, 0xF4, 0xEA)
AMBER_T = RGBColor(0xFE, 0xF7, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

X0, X1 = 0.35, 9.65      # content margins (in)
Y0, Y1 = 1.38, 5.42      # content top / bottom under the template title


# ----------------------------------------------------------------------------------------------- helpers
def _set_bullet(p, level_indent=0.16):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(level_indent))))
    pPr.set("indent", str(-int(Inches(level_indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar"))
    bc.set("char", "•")


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=3, name=None, margin=0.0):
    """paras: list of str | dict(text|runs, bullet, size, bold, color, align, space_after)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for para in paras:
        if isinstance(para, str):
            para = {"text": para}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", align)
        p.space_after = Pt(para.get("space_after", space_after))
        runs = para.get("runs") or [{"text": para.get("text", "")}]
        for r in runs:
            run = p.add_run()
            run.text = r["text"]
            f = run.font
            f.name = FONT
            f.size = Pt(r.get("size", para.get("size", size)))
            f.bold = r.get("bold", para.get("bold", bold))
            f.italic = r.get("italic", False)
            f.color.rgb = r.get("color", para.get("color", color))
            if r.get("link"):
                run.hyperlink.address = r["link"]
        if para.get("bullet"):
            _set_bullet(p)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def badge(slide, x, y, d, label, fill=BLUE, size=11):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return c


def arrow(slide, x, y, w=0.28, h=0.22, color=MUTED, shape=MSO_SHAPE.RIGHT_ARROW):
    a = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def picture(slide, path, x, y, w=None, h=None, border=True, name=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   width=Inches(w) if w else None, height=Inches(h) if h else None)
    if name:
        pic.name = name
    if border:
        pic.line.color.rgb = RGBColor(0xDA, 0xDC, 0xE0)
        pic.line.width = Pt(0.75)
    return pic


def phone_width(h):
    return h * 364 / 820


def table(slide, x, y, w, rows, col_w, size=10, header_fill=RGBColor(0x0B, 0x1F, 0x3A), row_h=0.28, bold_first_col=False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = str(val)
            run.font.name = FONT
            run.font.size = Pt(size)
            if r == 0:
                run.font.bold = True
                run.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            else:
                run.font.color.rgb = INK
                run.font.bold = bool(bold_first_col and c == 0)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else CARD
    return shp


def append_value(paragraph_el, value, bold=False, sep=" "):
    """Append ' value' as a run that copies the formatting of the paragraph's last run."""
    runs = paragraph_el.findall(qn("a:r"))
    src = runs[-1]
    new = copy.deepcopy(src)
    new.find(qn("a:t")).text = sep + value
    rpr = new.find(qn("a:rPr"))
    if rpr is not None:
        rpr.set("b", "1" if bold else "0")
    src.addnext(new)
    return new


def delete_slide(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    sld = list(sld_id_lst)[index]
    prs.part.drop_rel(sld.get(qn("r:id")))
    sld_id_lst.remove(sld)


def find_shape(slide, contains):
    for sh in slide.shapes:
        if sh.has_text_frame and contains in sh.text_frame.text:
            return sh
    raise KeyError(contains)


# ----------------------------------------------------------------------------------------------- build
prs = Presentation(str(HERE / "template.pptx"))
S = list(prs.slides)
(s_team, s_note, s_brief, s_opp, s_feat, s_flow, s_wire, s_arch, s_tech, s_cost, s_snap, s_perf,
 s_future, s_links, s_extra, s_thanks) = S

# ---- 1. Team details
tb = find_shape(s_team, "Team name")
for p in tb.text_frame.paragraphs:
    t = p.text.strip()
    if t.startswith("Team name"):
        append_value(p._p, "Prod Pilots", bold=True)
    elif t.startswith("Team leader name"):
        append_value(p._p, "Lohith DS", bold=True)
    elif t.startswith("Problem Statement"):
        append_value(p._p, "BFSI: Intelligent Risk, Fraud and Financial Experiences. Rakshak, a scam-moment guardian that stops authorised-push-payment scams at the moment of payment.", bold=False)

# ---- 3. Brief about the idea
text(s_brief, X0, Y0, 5.75, 0.62, [
    {"runs": [{"text": "Rakshak: ", "bold": True, "color": BLUE}, {"text": "stop the scam at the moment of payment, not after it", "bold": True}], "size": 15}], name="tagline")
text(s_brief, X0, Y0 + 0.62, 5.75, 2.0, [
    "Scam victims now send the money themselves, on a call, under pressure. Bank fraud engines see a genuine customer on a genuine device, and an SMS alert arrives after the money has left.",
    "Rakshak is a guardian built into a bank or UPI app. With the customer's opt-in it listens to a call from an unknown number and scores the caller's coercion live with Gemini. When the customer taps Pay, it fuses that call risk with the payee's mule-account risk from the bank's ledger, and in under 100 ms pauses the payment, explains why in the customer's language, alerts a trusted contact and files an auditable case. An ADK agent confirms with evidence; the customer always keeps an override.",
], size=11, space_after=6)
text(s_brief, X0, Y0 + 2.72, 5.75, 1.3, [
    {"text": "Built and deployed", "bold": True, "size": 12, "color": BLUE},
    {"text": "Customer app (installable PWA), caller phone, family phone and fraud desk on Cloud Run", "bullet": True},
    {"text": "Gemini 3.5 and 3.8 models, Google ADK 2.11 agents, Firestore case store", "bullet": True},
    {"text": "Live two-way calls between two phones, six languages, 50-call evaluation", "bullet": True},
], size=11, space_after=2)
stats = [("81.8%", "of Singapore's 2025 scam cases were transfers the victim made (Singapore Police)", RED),
         ("₹22,495 cr", "lost to cyber fraud in India in 2025, about 24% recovered (MHA via Moneylife)", RED),
         ("₹1.56 lakh", "average loss per digital-arrest victim in 2024 (Lok Sabha data)", AMBER)]
for i, (big, small, col) in enumerate(stats):
    y = Y0 + i * 1.36
    box(s_brief, 6.35, y, 3.3, 1.22, fill=CARD, name=f"stat{i}")
    text(s_brief, 6.5, y + 0.1, 3.0, 0.55, [{"text": big, "bold": True, "size": 26, "color": col}])
    text(s_brief, 6.5, y + 0.66, 3.0, 0.52, [{"text": small, "size": 10, "color": MUTED}])

# ---- 4. Opportunities: keep the heading, answer the three template questions as cards
tb = find_shape(s_opp, "Opportunities")
paras = tb.text_frame._txBody.findall(qn("a:p"))
for p in paras[1:]:
    p.getparent().remove(p)
first_p = paras[0]
for br in first_p.findall(qn("a:br")):
    first_p.remove(br)
for r in first_p.findall(qn("a:r")):
    t = r.find(qn("a:t"))
    t.text = t.text.replace("\x0b", "").replace("\v", "")
tb.height = Inches(0.61)
cards = [
    ("How is it different from existing ideas?", BLUE, BLUE_T, [
        "Pixel scam detection: on the handset, warns only, knows nothing about the payee",
        "Call-state signals (FICO Scam Signal, BioCatch): know a call is on, not what is said",
        "Bank fraud engines: blind to the call; the coerced victim passes every check",
        "SMS alerts and 1930: after the money left; about 24% is recovered",
        "Rakshak fuses call semantics with payee mule risk at the moment of authorisation, on the bank side",
    ]),
    ("How does it solve the problem?", RED, RED_T, [
        "Listens (opt-in) from the first ring; scores 8 coercion signals live in 6 languages",
        "Whispers privately to the customer when the risk passes 85%",
        "At Pay: graded decision in under 100 ms: proceed, warn, 30-minute hold, or hold plus alert",
        "Speaks the reason, notifies a trusted contact, drafts the fraud-desk case and 1930 referral",
        "Payments within 60 minutes of a high-risk call are still checked",
    ]),
    ("USP of the solution", GREEN, GREEN_T, [
        "Intervention, not an alert: the money does not leave until the customer has had time to think",
        "Explainable: every hold names its evidence; an ADK agent can only make it stricter",
        "Human in the loop: the customer can override by typing the amount",
        "Vernacular voice for the people scammers target most",
        "Privacy tiers: works with no audio, on-device text, or cloud analysis",
    ]),
]
cw = (X1 - X0 - 0.4) / 3
for i, (q, col, tint, items) in enumerate(cards):
    x = X0 + i * (cw + 0.2)
    box(s_opp, x, Y0, cw, Y1 - Y0, fill=tint, name=f"opp{i}")
    text(s_opp, x + 0.15, Y0 + 0.12, cw - 0.3, 0.5, [{"text": q, "bold": True, "size": 12, "color": col}])
    text(s_opp, x + 0.15, Y0 + 0.62, cw - 0.3, Y1 - Y0 - 0.75,
         [{"text": it, "bullet": True} for it in items], size=11, space_after=5)

# ---- 5. Features (3 x 3 grid)
feats = [
    ("Live coercion scoring", "Gemini scores 8 signals (authority, threat, secrecy, urgency, remote app, safe account, investment pressure, OTP) after every sentence"),
    ("Private whisper", "At 85% risk Rakshak warns the customer in their ear; the caller hears nothing"),
    ("Payment-moment hold", "Deterministic fusion of call risk, payee risk and profile decides in under 100 ms"),
    ("Voice barge-in", "The reason is spoken in the customer's language; Hindi, Hinglish, English, Kannada, Tamil, Singapore English"),
    ("Payee mule risk with lineage", "Fan-in, cash-out, dormancy, account age, telecom risk and 1930 complaints, computed from a live ledger stream"),
    ("ADK investigator", "Two Gemini agents check the payee and official advisories (RBI, CBI, MAS, SEBI) and confirm or upgrade"),
    ("Trusted-contact alert", "The family phone chimes with the case; also Pub/Sub and SMS hooks"),
    ("Talk to Rakshak", "Gemini Live voice conversation after the hold, with tools to keep the hold or call family"),
    ("Fraud desk", "Case timeline, agent trace, drafted 1930 referral, and the SMS-versus-Rakshak contrast"),
]
gw, gh = (X1 - X0 - 0.4) / 3, (Y1 - Y0 - 0.66) / 3
for i, (title_, desc) in enumerate(feats):
    r, c = divmod(i, 3)
    x, y = X0 + c * (gw + 0.2), Y0 + r * (gh + 0.12)
    box(s_feat, x, y, gw, gh, fill=CARD, name=f"feat{i}")
    badge(s_feat, x + 0.12, y + 0.12, 0.34, str(i + 1), fill=[BLUE, RED, GREEN][c])
    text(s_feat, x + 0.55, y + 0.1, gw - 0.65, 0.3, [{"text": title_, "bold": True, "size": 11}])
    text(s_feat, x + 0.55, y + 0.38, gw - 0.65, gh - 0.42, [{"text": desc, "size": 9, "color": MUTED}])
text(s_feat, X0, Y1 - 0.3, X1 - X0, 0.3, [{"runs": [
    {"text": "Also: ", "bold": True},
    {"text": "post-call carry-over (60 minutes), practice drill for elderly users, PIN-hidden balance, live two-way call between phones, installable app."}]}],
    size=10, color=MUTED)

# ---- 6. Process flow
steps = [
    [("1", "Caller rings", "unknown number; Guardian mode on (opt-in)"),
     ("2", "Live transcription", "Gemini 3.5 Transcribe Live, one stream per speaker"),
     ("3", "Coercion rubric", "Gemini 3.5 Flash, JSON, 8 signals, every sentence"),
     ("4", "Private whisper", "at 85% risk, Gemini 3.8 Flash TTS")],
    [("5", "Customer taps Pay", "during the call or within 60 min after"),
     ("6", "Fusion policy", "0.65 call + 0.35 payee + profile, under 100 ms"),
     ("7", "Graded action", "proceed, warn, 30-min hold, hold + alert"),
     ("8", "ADK agents confirm", "payee ledger + official advisories; never relax")],
    [("9", "Voice + family alert", "reason spoken; trusted contact notified"),
     ("10", "Customer choice", "cancel, Talk to Rakshak (Gemini Live), or typed override"),
     ("11", "Fraud desk", "case in Firestore, agent trace, drafted 1930 referral"),
     ("12", "Bank acts", "freeze review of the mule account; customer follow-up")],
]
phase = [("During the call", BLUE, BLUE_T), ("At the moment of payment", RED, RED_T), ("After the decision", GREEN, GREEN_T)]
bw, bh = 1.88, 0.86
for r, row in enumerate(steps):
    y = Y0 + 0.05 + r * 1.33
    label, col, tint = phase[r]
    text(s_flow, X0, y - 0.02, 1.1, bh, [{"text": label, "bold": True, "size": 10, "color": col}], anchor=MSO_ANCHOR.MIDDLE)
    for c, (n, head, sub) in enumerate(row):
        x = X0 + 1.15 + c * (bw + 0.27)
        box(s_flow, x, y, bw, bh, fill=tint, name=f"flow{r}{c}")
        badge(s_flow, x + 0.08, y + 0.08, 0.3, n, fill=col, size=9)
        text(s_flow, x + 0.44, y + 0.07, bw - 0.5, 0.3, [{"text": head, "bold": True, "size": 10}])
        text(s_flow, x + 0.1, y + 0.4, bw - 0.18, bh - 0.42, [{"text": sub, "size": 8.5, "color": MUTED}])
        if c < 3:
            arrow(s_flow, x + bw + 0.02, y + bh / 2 - 0.1, w=0.23, h=0.2, color=col)
    if r < 2:
        arrow(s_flow, X0 + 1.15 + 3 * (bw + 0.27) + bw / 2 - 0.1, y + bh + 0.06, w=0.2, h=0.3, color=MUTED, shape=MSO_SHAPE.DOWN_ARROW)

# ---- 7. Wireframes: the customer journey as real screens
journey = [("p_home", "Home: Guardian on, balance hidden by PIN"), ("p_incoming", "Unknown number rings"),
           ("p_risk", "Live risk 100%, whisper, reasons"), ("p_pay", "Pay screen knows a call is on"),
           ("p_hold", "Payment paused, reason spoken")]
ph = 3.45
pw = phone_width(ph)
gap = (X1 - X0 - 5 * pw) / 4
for i, (img, cap) in enumerate(journey):
    x = X0 + i * (pw + gap)
    picture(s_wire, SHOTS / f"{img}.png", x, Y0, h=ph, name=f"wire{i}")
    text(s_wire, x - 0.1, Y0 + ph + 0.05, pw + 0.2, 0.45, [{"text": cap, "size": 9, "color": MUTED, "align": PP_ALIGN.CENTER}])
    if i < 4:
        arrow(s_wire, x + pw + gap / 2 - 0.12, Y0 + ph / 2 - 0.1, w=0.24, h=0.2, color=BLUE)

# ---- 8. Architecture (rendered from the system diagram) + one line on the Google Cloud services
d = pdfium.PdfDocument(str(HERE.parent / "Rakshak_Deck.pdf"))
img = d[5].render(scale=2.0).to_pil().convert("RGB")
W, H = img.size
img.crop((0, int(H * 0.198), W, int(H * 0.893))).save(SHOTS / "architecture_hi.png")
aw = X1 - X0
ah = aw * (int(H * 0.893) - int(H * 0.198)) / W
picture(s_arch, SHOTS / "architecture_hi.png", X0, Y0 - 0.02, w=aw, border=False, name="architecture")
hosting_bit = f" Firebase Hosting serves the same app at {HOSTING}." if HOSTING else ""
text(s_arch, X0, Y0 - 0.02 + ah + 0.06, aw, 0.4, [{"runs": [
    {"text": "Google Cloud: ", "bold": True, "color": BLUE},
    {"text": "Cloud Run (FastAPI, WebSockets, one always-on instance) · Gemini on Vertex AI and the Gemini API · ADK 2.11 · Firestore case store (Firebase) · Pub/Sub alerts · BigQuery ledger SQL · Secret Manager." + hosting_bit}]}],
    size=9, color=MUTED)

# ---- 9. Technologies (six cards)
tech = [
    ("Gemini models", BLUE, ["3.5 Flash: coercion rubric (JSON), ADK agents, case narrative", "3.5 Transcribe Live: one live stream per speaker", "3.8 Flash TTS: whisper and spoken hold", "3.8 Live: voice conversation with tools", "3.8 Flash: agent fallback"]),
    ("Agents", RED, ["Google ADK 2.11 SequentialAgent", "risk_investigator: payee ledger + official advisories, Search grounding fallback", "decision_agent: place_payment_hold tool", "Policy floor: agents can only make it stricter"]),
    ("Serving", GREEN, ["Cloud Run, asia-southeast1, WebSockets, session affinity", "FastAPI, Python 3.12, uv", "Installable PWA: customer app, caller phone, family phone, fraud desk"] + (["Firebase Hosting front door"] if HOSTING else [])),
    ("Data", BLUE, ["Firestore (Firebase): cases, profiles, alerts", "Ledger stream with mule features (same formula as the BigQuery SQL)", "Pub/Sub for trusted-contact alerts; BigQuery sink optional"]),
    ("Security and privacy", RED, ["Secret Manager for the API key; IAM service account", "Opt-in Guardian mode, visible indicator", "No raw audio stored; PIN-hidden balance", "Typed-amount override, holds expire"]),
    ("Quality", GREEN, ["50-call labelled evaluation in 6 languages", "16 automated tests incl. a 3-day ledger drift test", "Live two-device audio smoke test on the public URL"]),
]
tw, th = (X1 - X0 - 0.4) / 3, (Y1 - Y0 - 0.15) / 2
for i, (head, col, items) in enumerate(tech):
    r, c = divmod(i, 3)
    x, y = X0 + c * (tw + 0.2), Y0 + r * (th + 0.15)
    box(s_tech, x, y, tw, th, fill=CARD, name=f"tech{i}")
    text(s_tech, x + 0.14, y + 0.08, tw - 0.28, 0.3, [{"text": head, "bold": True, "size": 12, "color": col}])
    text(s_tech, x + 0.14, y + 0.42, tw - 0.28, th - 0.48, [{"text": it, "bullet": True} for it in items], size=9.5, space_after=2)

# ---- 10. Estimated implementation cost
rows = [["Item (per guarded 3-minute call)", "Assumption", "Estimate (USD)"],
        ["Live transcription, both speakers", "Gemini 3.5 Transcribe Live, 360 s of audio at 25 tokens/s", "0.03"],
        ["Coercion rubric", "about 15 scorings, Gemini 3.5 Flash", "0.07"],
        ["ADK agents and case narrative", "about 5 Gemini 3.5 Flash calls", "0.03"],
        ["Whisper and spoken hold", "Gemini 3.8 Flash TTS, about 33 s of audio", "0.01"],
        ["Total per guarded call with a hold", "list prices, Gemini API, Oct 2026", "about 0.14 (₹12)"]]
table(s_cost, X0, Y0, 5.9, rows, [2.0, 2.85, 1.05], size=9, row_h=0.42)
rows2 = [["Infrastructure (prototype)", "Estimate"],
         ["Cloud Run, 1 always-on instance (1 vCPU, 1 GiB)", "about USD 50 to 70 a month"],
         ["Firestore, Pub/Sub, BigQuery, Secret Manager", "within free tiers at prototype scale"]]
table(s_cost, X0, Y0 + 2.7, 5.9, rows2, [3.6, 2.3], size=9, row_h=0.38)
box(s_cost, 6.5, Y0, X1 - 6.5, 2.4, fill=GREEN_T, name="costcallout")
text(s_cost, 6.65, Y0 + 0.12, X1 - 6.8, 2.2, [
    {"text": "₹12", "bold": True, "size": 30, "color": GREEN},
    {"text": "per guarded call that ends in a hold", "size": 11},
    {"text": "versus ₹1.56 lakh lost by the average digital-arrest victim", "size": 11, "bold": True, "space_after": 6},
    {"text": "Cost is paid only when Guardian mode is on and an unknown number calls.", "size": 9, "color": MUTED}])
box(s_cost, 6.5, Y0 + 2.55, X1 - 6.5, 1.45, fill=CARD, name="costnotes")
text(s_cost, 6.65, Y0 + 2.62, X1 - 6.8, 1.35, [
    {"text": "At bank scale", "bold": True, "size": 11, "color": BLUE},
    {"text": "Sessions move to Firestore or Memorystore and Cloud Run scales out; agents can move to Agent Runtime.", "bullet": True},
    {"text": "Estimates use list prices; Vertex AI contract prices differ.", "bullet": True}], size=9)

# ---- 11. Snapshots of the prototype
snaps = [("p_caller", "Caller phone: live two-way call"), ("p_payee", "Payee evidence with ledger lineage"), ("p_family", "Family phone: alert arrives")]
sh_h = 3.5
sw = phone_width(sh_h)
for i, (img, cap) in enumerate(snaps):
    x = X0 + i * (sw + 0.25)
    picture(s_snap, SHOTS / f"{img}.png", x, Y0, h=sh_h, name=f"snap{i}")
    text(s_snap, x - 0.1, Y0 + sh_h + 0.05, sw + 0.2, 0.45, [{"text": cap, "size": 9, "color": MUTED, "align": PP_ALIGN.CENTER}])
dx = X0 + 3 * (sw + 0.25) + 0.05
dw = X1 - dx
dh = dw * 1044 / 972
if dh > sh_h:
    dh = sh_h
    dw = dh * 972 / 1044
picture(s_snap, SHOTS / "desk_case.png", dx, Y0, w=dw, name="snapdesk")
text(s_snap, dx, Y0 + dh + 0.05, dw, 0.45, [{"text": "Fraud desk: case, timeline, agent versus policy, SMS contrast", "size": 9, "color": MUTED, "align": PP_ALIGN.CENTER}])

# ---- 12. Performance report
tiles = [("1.00", "Precision", "no benign call held"), ("1.00", "Recall", "every scam call held"),
         ("39 / 39", "Payment decisions", "correct action on every payment"), ("93.2%", "Signal accuracy", "8 signals x 50 calls")]
tw2, th2 = 2.05, 1.05
for i, (big, lab, sub) in enumerate(tiles):
    r, c = divmod(i, 2)
    x, y = X0 + c * (tw2 + 0.15), Y0 + r * (th2 + 0.15)
    box(s_perf, x, y, tw2, th2, fill=GREEN_T if i < 3 else BLUE_T, name=f"perf{i}")
    text(s_perf, x + 0.12, y + 0.06, tw2 - 0.2, 0.5, [{"text": big, "bold": True, "size": 22, "color": GREEN if i < 3 else BLUE}])
    text(s_perf, x + 0.12, y + 0.56, tw2 - 0.2, 0.45, [{"text": lab, "bold": True, "size": 10}, {"text": sub, "size": 8.5, "color": MUTED}], space_after=0)
cm = [["50 labelled calls", "Predicted scam", "Predicted benign"], ["Actual scam (25)", "25", "0"], ["Actual benign (25)", "0", "25"]]
table(s_perf, X0, Y0 + 2.48, 4.25, cm, [1.65, 1.3, 1.3], size=9.5, row_h=0.32, bold_first_col=True)
text(s_perf, X0, Y0 + 3.5, 4.25, 0.55, [{"text": "25 scam and 25 benign calls in Hindi, Hinglish, English, Singapore English, Kannada and Tamil, including 10 legitimate-but-urgent calls and 10 scams written to dodge keywords. Self-authored set, not held-out.", "size": 8.5, "color": MUTED}])
lat = [["Latency and reliability", "Measured"],
       ["Coercion score after a sentence", "3.1 s mean, 2.0 s median"],
       ["Hold on screen after Pay", "0.02 to 1.4 s"],
       ["Spoken reason (prewarmed)", "about 0.1 s after the hold"],
       ["ADK agent confirmation", "5.4 to 7.2 s"],
       ["Live two-device audio test", "3 of 3 sentences transcribed; whisper fired"],
       ["Automated tests", "16 passing, incl. 3-day ledger drift"],
       ["Ledger stream", "about 12,800 transactions in window, about 120 a minute"]]
table(s_perf, 4.85, Y0, X1 - 4.85, lat, [2.35, X1 - 4.85 - 2.35], size=9, row_h=0.43, bold_first_col=True)

# ---- 13. Additional details / future development
box(s_future, X0, Y0, 4.55, Y1 - Y0, fill=BLUE_T, name="roadmap")
text(s_future, X0 + 0.15, Y0 + 0.1, 4.25, Y1 - Y0 - 0.2, [
    {"text": "Next with a bank or PSP partner", "bold": True, "size": 12, "color": BLUE, "space_after": 6},
    {"text": "Real ledger features in BigQuery on a schedule; DoT FRI, NPCI DPIP and I4C registry feeds", "bullet": True},
    {"text": "On-device Tier 1 with an OS or OEM partner (Pixel scam detection to a GPay hold)", "bullet": True},
    {"text": "Telco audio fork for cellular calls, as banks already do with call-state signals", "bullet": True},
    {"text": "Native Android app: call-state signal, foreground service, FCM push to family", "bullet": True},
    {"text": "Human-labelled call corpus and red-team evaluation; native-speaker QA for Indic voices", "bullet": True},
    {"text": "Model-risk evidence pack for RBI FREE-AI and MAS AI risk guidelines", "bullet": True},
], size=11.5, space_after=6)
box(s_future, X0 + 4.75, Y0, X1 - X0 - 4.75, Y1 - Y0, fill=AMBER_T, name="limits")
text(s_future, X0 + 4.9, Y0 + 0.1, X1 - X0 - 5.05, Y1 - Y0 - 0.2, [
    {"text": "Honest limits of this prototype", "bold": True, "size": 12, "color": AMBER, "space_after": 6},
    {"text": "Synthetic ledger and scripted calls; no bank partner yet", "bullet": True},
    {"text": "Third-party apps cannot tap cellular calls on Android or iOS, so live audio uses in-app calls or a telco bridge", "bullet": True},
    {"text": "The Twilio phone bridge is one-way: the spoken hold stays on the app", "bullet": True},
    {"text": "Kannada and Tamil voices not yet tested with native speakers", "bullet": True},
    {"text": "One Cloud Run instance holds live call sessions; scale-out needs shared session state", "bullet": True},
], size=11.5, space_after=6)

# ---- 14. Links
tb = find_shape(s_links, "Provide links")
for p in tb.text_frame.paragraphs:
    t = p.text.strip()
    if t.startswith("GitHub Public Repository"):
        append_value(p._p, REPO, bold=False, sep=": ")
    elif t.startswith("Demo Video Link"):
        append_value(p._p, "[paste the public YouTube link here]", bold=True, sep=": ")
    elif t.startswith("Final Product Link"):
        append_value(p._p, URL, bold=False, sep=": ")
box(s_links, X0, 2.4, X1 - X0, 2.15, fill=CARD, name="howtotest")
text(s_links, X0 + 0.18, 2.5, (X1 - X0) / 2 - 0.3, 1.95, [
    {"text": "Try it in 60 seconds (one device)", "bold": True, "size": 12, "color": BLUE, "space_after": 6},
    {"text": "Open the product link, tap Test an incoming call", "bullet": True},
    {"text": "Choose Cyber Crime Branch (English), tap Auto-play script", "bullet": True},
    {"text": "When asked to pay, tap Open UPI pay, then Pay", "bullet": True},
    {"text": "Watch the hold, hear Rakshak, open /desk for the case", "bullet": True},
    {"text": "Balance PIN for the demo: 1234", "bullet": True},
], size=10, space_after=3)
mid = X0 + (X1 - X0) / 2
text(s_links, mid, 2.5, (X1 - X0) / 2 - 0.2, 1.95, [
    {"text": "Live two-way call (two phones)", "bold": True, "size": 12, "color": RED, "space_after": 6},
    {"text": "Customer phone: /?session=demo1&user=demo-user", "bullet": True},
    {"text": "Caller phone: /call?session=demo1, tap Call", "bullet": True},
    {"text": "Family phone: /family?user=demo-user", "bullet": True},
    {"text": "Fraud desk: /desk", "bullet": True},
    {"text": "Hindi customer: &user=demo-user-hi", "bullet": True},
], size=10, space_after=3)

# ---- 15 (extra template slide). Privacy by design and real cellular calls
title_src = find_shape(s_brief, "Brief about the idea")
new_title = copy.deepcopy(title_src._element)
s_extra.shapes._spTree.append(new_title)
for t in new_title.iter(qn("a:t")):
    t.text = ""
first_t = next(new_title.iter(qn("a:t")))
first_t.text = "Privacy by design, and how it works on a real cellular call"
rows3 = [["Tier", "What Rakshak uses", "What leaves the phone"],
         ["0: no audio", "Call from an unknown number + pay screen + payee risk", "Only the payment request"],
         ["1: on-device text", "Rolling 60 s transcript kept on the phone", "Scores and last 60 s, only if a payment is attempted"],
         ["2: cloud (this prototype)", "Live transcription through a lawful channel", "Audio stream, never stored"]]
table(s_extra, X0, Y0, 5.6, rows3, [1.45, 2.25, 1.9], size=9, row_h=0.5, bold_first_col=True)
text(s_extra, X0, Y0 + 2.15, 5.6, 1.8, [
    {"text": "Rules at every tier", "bold": True, "size": 11, "color": BLUE},
    {"text": "Guardian mode is opt-in with a visible listening indicator", "bullet": True},
    {"text": "The caller never hears Rakshak; holds expire; the customer can override", "bullet": True},
    {"text": "Raw audio is never written to disk; cases are kept for the bank's audit trail", "bullet": True},
], size=10, space_after=3)
box(s_extra, 6.2, Y0, X1 - 6.2, Y1 - Y0, fill=BLUE_T, name="cellular")
text(s_extra, 6.35, Y0 + 0.1, X1 - 6.5, Y1 - Y0 - 0.2, [
    {"text": "Real cellular calls", "bold": True, "size": 12, "color": BLUE, "space_after": 6},
    {"text": "No third-party app can read cellular call audio on Android 10+ or iOS. Rakshak therefore runs on the bank side and receives audio only through lawful channels:", "size": 10, "space_after": 6},
    {"text": "the bank's own in-app or verified calls (shown in the demo)", "bullet": True},
    {"text": "a telco audio fork with consent (FICO Scam Signal precedent)", "bullet": True},
    {"text": "an OS vendor feed: Google can connect Pixel scam detection to a GPay hold", "bullet": True},
    {"text": "Without audio, Tier 0 still holds risky payments to mule accounts during unknown calls.", "size": 10, "bold": True},
], size=10, space_after=3)

# ---- remove the template's instruction slide
delete_slide(prs, 1)

prs.core_properties.title = "Rakshak by Prod Pilots"
prs.core_properties.author = "Prod Pilots (Lohith DS)"
prs.save(str(OUT))
print("wrote", OUT)
